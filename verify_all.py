# -*- coding: utf-8 -*-
"""Независимая сверка всех Excel с исходниками: КАРТА, количество строк,
словарь, дубли внутри файлов и между файлами.
python verify_all.py
"""
import sys
from openpyxl import load_workbook
from promptlib import check_rows, duration, similar, source_lines

PAIRS = [
    ("Gaia_planety.xlsx", "src/src_gaia.txt"),
    ("Budushchee_Zemli.xlsx", "src/src_future.txt"),
    ("Zvezda_chernoy_dyry.xlsx", "src/src_bhstar.txt"),
    ("Stupeni_na_Lune.xlsx", "src/src_moonstage.txt"),
    ("Teoriya_strun.xlsx", "src/src_strings.txt"),
    ("Panspermiya.xlsx", "src/src_pansp.txt"),
    ("Neptun.xlsx", "src/src_neptune.txt"),
    ("Planety_okeany.xlsx", "src/src_oceanw.txt"),
    ("Proshloe_Zemli.xlsx", "src/src_pastearth.txt"),
    ("Graviton.xlsx", "src/src_graviton.txt"),
    ("Lovushka_gravitona.xlsx", "src/src_gravcatch.txt"),
    ("Mashina_vremeni_Gedelya.xlsx", "src/src_goedel.txt"),
    ("Temnaya_materiya_CERN.xlsx", "src/src_darkmatter.txt"),
    ("Lestnitsa_chernyh_dyr.xlsx", "src/src_bhladder.txt"),
    ("Korabl_bystree_sveta.xlsx", "src/src_ftlship.txt"),
]

total_err = 0
all_nano = []
print(f"{'файл':28} {'кадров':>6} {'время':>6} {'космос':>6}  ошибок")
for xlsx, txt in PAIRS:
    try:
        ws = load_workbook(xlsx)["Промты"]
    except FileNotFoundError:
        print(f"{xlsx:28} НЕТ ФАЙЛА")
        total_err += 1
        continue
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        rows.append({"part": r[1], "space": r[2] == "да", "nano": r[3], "veo": r[4], "karta": r[5]})
    src = source_lines(txt)
    errs, _ = check_rows(rows, src)
    share = sum(r["space"] for r in rows) / max(1, len(rows))
    print(f"{xlsx:28} {len(rows):>6} {duration(src):>6} {share:>6.0%}  {len(errs)}")
    for e in errs:
        print("   ERR", e)
    total_err += len(errs)
    all_nano += [(xlsx, i, r["nano"]) for i, r in enumerate(rows, 1)]

cross = 0
for a in range(len(all_nano)):
    for b in range(a + 1, len(all_nano)):
        fa, ia, na = all_nano[a]
        fb, ib, nb = all_nano[b]
        if fa == fb:
            continue
        s = similar(na, nb)
        if s >= 0.8:
            cross += 1
            print(f"   ERR между файлами: {fa}#{ia} ~ {fb}#{ib} ({s:.0%})")
total_err += cross
print("ИТОГ:", "чисто" if total_err == 0 else f"{total_err} ошибок")
sys.exit(1 if total_err else 0)
