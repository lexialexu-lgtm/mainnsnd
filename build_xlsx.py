# -*- coding: utf-8 -*-
"""Проверяет JSON и собирает стилизованный Excel.
python build_xlsx.py rows.json out.xlsx source.txt
Если есть ошибки — Excel не пишется, код выхода 1.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from promptlib import check_rows, duration, load, source_lines, space_share

rows_path, out_path, src_path = sys.argv[1], sys.argv[2], sys.argv[3]
rows = load(rows_path)
src = source_lines(src_path)
errs, warns = check_rows(rows, src)
for w in warns:
    print("WARN", w)
if errs:
    for e in errs:
        print("ERR ", e)
    sys.exit(f"ПРОВЕРКА НЕ ПРОЙДЕНА: {len(errs)} ошибок, Excel не собран")

wb = Workbook()
ws = wb.active
ws.title = "Промты"
head = ["№", "Часть", "Космос", "NANO", "VEO", "КАРТА"]
ws.append(head)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfill = PatternFill("solid", fgColor="1F2A44")
part_fills = ["FFFFFF", "F3F6FB", "FFFFFF", "F3F6FB", "FFFFFF"]
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = hfill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border
for i, r in enumerate(rows, 1):
    ws.append([i, r["part"], "да" if r["space"] else "нет", r["nano"], r["veo"], r["karta"]])
    fill = PatternFill("solid", fgColor=part_fills[(r["part"] - 1) % 5])
    for c in ws[i + 1]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = border
        c.fill = fill
    ws.cell(i + 1, 3).font = Font(color="1F5FBF" if r["space"] else "7F7F7F")
for col, w in zip("ABCDEF", [6, 7, 8, 80, 26, 60]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:F{len(rows) + 1}"

s = wb.create_sheet("Сводка")
for row in [("Кадров", len(rows)), ("Хронометраж (оценка)", duration(src)),
            ("Доля космоса", f"{space_share(rows):.0%}")]:
    s.append(row)
for c in s["A"]:
    c.font = Font(bold=True)
s.column_dimensions["A"].width = 24
s.column_dimensions["B"].width = 14
wb.save(out_path)
print(f"OK: {out_path} — {len(rows)} кадров, {duration(src)}, космос {space_share(rows):.0%}, предупреждений {len(warns)}")
